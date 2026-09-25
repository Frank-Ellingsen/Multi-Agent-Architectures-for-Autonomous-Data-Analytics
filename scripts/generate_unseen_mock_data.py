"""Generate authentic unseen mock datasets across multiple businesses and formats.

Formats:
- Excel (.xlsx) with multiple sheets
- Tab-delimited text (.tsv)
- Semicolon-delimited European CSV (.csv) with comma decimals
- Pipe-delimited text (.psv)
- Formatted PDF (.pdf) with structured tabular data
- Standard comma-delimited text (.txt)

Businesses:
1. Maritime & Naval Defense Engineering Project Controlling (Umoe Mandal / naval composite shipbuilder)
2. Cloud SaaS & Subscription Business
3. Nordic Retail & Supply Chain
4. Offshore Energy & Marine Drilling Operations
5. Healthcare & Regional Hospital Operations
6. Professional Consulting & Engineering Advisory
"""

from __future__ import annotations

import csv
from pathlib import Path
import openpyxl
from reportlab.lib.pagesizes import letter, landscape
from reportlab.lib import colors
from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle


def generate_maritime_defense_excel(output_path: Path) -> Path:
    wb = openpyxl.Workbook()
    
    # Sheet 1: Project_Summary
    ws_summary = wb.active
    ws_summary.title = "Project_Summary"
    ws_summary.append([
        "Project_ID", "Vessel_Class", "Customer", "Contract_Value_NOK", 
        "Baseline_Start", "Planned_Delivery", "Current_Phase", "Currency"
    ])
    ws_summary.append([
        "P-901-CORVETTE", "Fast Missile Patrol Vessel (Composite Sandwich)", 
        "Royal Norwegian Navy", 485000000.0, "2024-01-15", "2026-11-30", "Production & Outfitting", "NOK"
    ])
    ws_summary.append([
        "P-902-MINESWEEPER", "Mine Countermeasures Vessel (MCMV)", 
        "Naval Material Command", 320000000.0, "2024-06-01", "2027-04-15", "Hull Fabrication", "NOK"
    ])

    # Sheet 2: WBS_Cost_Control
    ws_wbs = wb.create_sheet(title="WBS_Cost_Control")
    ws_wbs.append([
        "WBS_Code", "WBS_Description", "Department", "Budget_NOK", 
        "Actual_Cost_NOK", "ETC_NOK", "EAC_NOK", "SPI", "CPI", "Hours_Actual"
    ])
    wbs_data = [
        ("WP-100", "Hull & Composite Superstructure", "Production Composite", 125000000.0, 132500000.0, 18000000.0, 150500000.0, 0.94, 0.88, 42500.0),
        ("WP-200", "Propulsion & Gas Turbine Integration", "Mechanical Engineering", 95000000.0, 89200000.0, 12000000.0, 101200000.0, 0.98, 0.96, 18200.0),
        ("WP-300", "Combat Management System & Radar", "Systems Integration", 145000000.0, 141800000.0, 15500000.0, 157300000.0, 0.99, 0.95, 24600.0),
        ("WP-400", "Electrical & Auxiliary Systems", "Electrical Engineering", 42000000.0, 38600000.0, 4500000.0, 43100000.0, 1.01, 0.97, 14300.0),
        ("WP-500", "Sea Acceptance Trials & Commissioning", "Project Management", 38000000.0, 14500000.0, 26500000.0, 41000000.0, 0.92, 0.91, 6800.0),
        ("WP-900", "Management & Technical Contingency", "Project Controlling", 40000000.0, 0.0, 30000000.0, 30000000.0, 1.00, 1.00, 1200.0),
    ]
    for row in wbs_data:
        ws_wbs.append(list(row))

    # Sheet 3: Milestone_Risk
    ws_risk = wb.create_sheet(title="Milestone_Risk")
    ws_risk.append(["Milestone_ID", "Milestone_Name", "Target_Date", "Status", "Risk_Level", "Contingency_Allocation_NOK"])
    risk_data = [
        ("M-01", "Composite Hull Moulding Complete", "2024-06-30", "Completed", "Low", 0.0),
        ("M-02", "Superstructure Infusion & Bonding", "2024-12-15", "Completed", "Medium", 2500000.0),
        ("M-03", "Main Gas Turbine Alignment", "2025-05-20", "Delayed", "High", 6000000.0),
        ("M-04", "Harbor Acceptance Test (HAT)", "2025-10-15", "In Progress", "High", 8500000.0),
        ("M-05", "Sea Acceptance Trials (SAT)", "2026-06-30", "Planned", "Critical", 12000000.0),
    ]
    for row in risk_data:
        ws_risk.append(list(row))

    output_path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(output_path)
    return output_path


def generate_saas_tsv(output_path: Path) -> Path:
    headers = [
        "customer_id", "company_name", "plan_tier", "mrr_usd", "arr_usd",
        "budget_mrr_usd", "forecast_mrr_usd", "churn_risk_score", "billing_cycle",
        "cac_usd", "ltv_usd", "active_seats"
    ]
    rows = [
        ["CUST-001", "Nordic Fintech AS", "Enterprise", "12500.00", "150000.00", "11000.00", "13000.00", "0.08", "Annual", "18500.00", "225000.00", "145"],
        ["CUST-002", "Arctic Logistics Group", "Enterprise", "9800.00", "117600.00", "9500.00", "10200.00", "0.14", "Annual", "14200.00", "176000.00", "98"],
        ["CUST-003", "Scandic Retail Tech", "Professional", "4500.00", "54000.00", "5000.00", "4800.00", "0.32", "Monthly", "7800.00", "65000.00", "45"],
        ["CUST-004", "Oslo Health Systems", "Enterprise", "18200.00", "218400.00", "16500.00", "19000.00", "0.05", "Annual", "24000.00", "350000.00", "210"],
        ["CUST-005", "Fjord Media Network", "Growth", "2800.00", "33600.00", "3200.00", "2900.00", "0.45", "Monthly", "5100.00", "38000.00", "28"],
        ["CUST-006", "Bergen Green Energy", "Enterprise", "14200.00", "170400.00", "13500.00", "15000.00", "0.11", "Annual", "19500.00", "240000.00", "160"],
        ["CUST-007", "Stavanger Subsea AI", "Professional", "6200.00", "74400.00", "6000.00", "6500.00", "0.18", "Annual", "9200.00", "110000.00", "62"],
        ["CUST-008", "Telemark Data Solutions", "Growth", "3400.00", "40800.00", "3800.00", "3500.00", "0.28", "Monthly", "5600.00", "48000.00", "35"],
        ["CUST-009", "Trondheim Quantum Labs", "Professional", "5800.00", "69600.00", "5500.00", "6000.00", "0.12", "Annual", "8800.00", "98000.00", "55"],
        ["CUST-010", "Drammen Mobility AS", "Starter", "1200.00", "14400.00", "1500.00", "1300.00", "0.52", "Monthly", "2800.00", "18000.00", "12"],
    ]
    output_path.parent.mkdir(parents=True, exist_ok=True)
    with output_path.open("w", newline="", encoding="utf-8") as f:
        writer = csv.writer(f, delimiter="\t")
        writer.writerow(headers)
        writer.writerows(rows)
    return output_path


def generate_nordic_retail_semicolon_csv(output_path: Path) -> Path:
    headers = [
        "Artikkel_ID", "Varenavn", "Kategori", "Lager_Lokasjon",
        "Innkjoepspris_NOK", "Salgspris_NOK", "Lagerbeholdning",
        "Maanedlig_Salg_NOK", "Budsjett_Salg_NOK", "Prognose_Salg_NOK", "Bruttomargin_Prosent"
    ]
    rows = [
        ["ART-101", "Vinterjakke Ekspedisjon", "Yttertoy", "Lager Oslo", "1250,50", "2899,00", "450", "1304550,00", "1200000,00", "1350000,00", "56,8"],
        ["ART-102", "Ullgenser Merino", "Strikk", "Lager Bergen", "420,00", "999,00", "820", "819180,00", "850000,00", "820000,00", "57,9"],
        ["ART-103", "Fjellstovel Gore-Tex", "Fottoy", "Lager Kristiansand", "980,00", "2299,00", "310", "712690,00", "750000,00", "720000,00", "57,3"],
        ["ART-104", "Teknisk Skallbukse", "Yttertoy", "Lager Oslo", "790,00", "1799,00", "540", "971460,00", "900000,00", "980000,00", "56,0"],
        ["ART-105", "Dunjakke Lettvekt", "Yttertoy", "Lager Trondheim", "890,50", "2199,00", "620", "1363380,00", "1400000,00", "1380000,00", "59,5"],
        ["ART-106", "Vandresekk 65L Pro", "Utstyr", "Lager Oslo", "650,00", "1599,00", "280", "447720,00", "420000,00", "450000,00", "59,3"],
        ["ART-107", "Termoundertoy Sett", "Ull", "Lager Bergen", "210,00", "599,00", "1450", "868550,00", "800000,00", "870000,00", "64,9"],
        ["ART-108", "Vinterhanske Vindtett", "Tilbehor", "Lager Kristiansand", "145,00", "399,00", "980", "391020,00", "350000,00", "400000,00", "63,6"],
    ]
    output_path.parent.mkdir(parents=True, exist_ok=True)
    with output_path.open("w", newline="", encoding="utf-8-sig") as f:
        writer = csv.writer(f, delimiter=";")
        writer.writerow(headers)
        writer.writerows(rows)
    return output_path


def generate_offshore_marine_psv(output_path: Path) -> Path:
    headers = [
        "vessel_id", "vessel_name", "operator", "day_rate_usd",
        "operating_hours", "downtime_hours", "fuel_cost_usd",
        "maintenance_actual_usd", "maintenance_budget_usd", "maintenance_forecast_usd", "crew_fte"
    ]
    rows = [
        ["VES-801", "Northern Pioneer", "Equinor Energy", "42500.00", "720.0", "12.5", "185000.00", "945000.00", "850000.00", "980000.00", "28.5"],
        ["VES-802", "Arctic Protector", "Aker BP", "51000.00", "708.0", "24.0", "210000.00", "1120000.00", "980000.00", "1150000.00", "34.0"],
        ["VES-803", "Fjord Constructor", "Subsea 7", "68000.00", "695.0", "38.5", "295000.00", "1650000.00", "1450000.00", "1700000.00", "46.0"],
        ["VES-804", "Viking Supporter", "ConocoPhillips", "36000.00", "735.0", "4.0", "145000.00", "680000.00", "720000.00", "700000.00", "22.0"],
        ["VES-805", "Skagerrak Explorer", "Vår Energi", "58000.00", "712.0", "18.0", "240000.00", "1380000.00", "1250000.00", "1400000.00", "38.5"],
    ]
    output_path.parent.mkdir(parents=True, exist_ok=True)
    with output_path.open("w", newline="", encoding="utf-8") as f:
        writer = csv.writer(f, delimiter="|")
        writer.writerow(headers)
        writer.writerows(rows)
    return output_path


def generate_hospital_kpi_pdf(output_path: Path) -> Path:
    output_path.parent.mkdir(parents=True, exist_ok=True)
    doc = SimpleDocTemplate(
        str(output_path),
        pagesize=landscape(letter),
        leftMargin=36,
        rightMargin=36,
        topMargin=36,
        bottomMargin=36
    )
    styles = getSampleStyleSheet()
    title_style = ParagraphStyle(
        'DocTitle',
        parent=styles['Heading1'],
        fontName='Helvetica-Bold',
        fontSize=18,
        leading=22,
        textColor=colors.HexColor('#0F172A'),
        spaceAfter=6
    )
    subtitle_style = ParagraphStyle(
        'DocSubtitle',
        parent=styles['Normal'],
        fontName='Helvetica',
        fontSize=10,
        textColor=colors.HexColor('#475569'),
        spaceAfter=14
    )

    story = []
    story.append(Paragraph("Southern Hospital Trust - Q3 Operational & Financial Review", title_style))
    story.append(Paragraph("Autonomous Multi-Agent Controlling & Clinical Resource Performance Report (Edward Tufte Low-Ink Table)", subtitle_style))
    story.append(Spacer(1, 10))

    headers = [
        "Department", "Admissions", "Bed_Days", "Occupancy_Pct", 
        "Budget_Cost_NOK", "Actual_Cost_NOK", "Forecast_Cost_NOK", "Variance_NOK", "Staff_FTE"
    ]
    table_data = [headers]
    rows = [
        ["Emergency & Trauma", "4250", "14200", "94.5%", "48500000.00", "52400000.00", "54000000.00", "3900000.00", "145.0"],
        ["Cardiology & Thoracic", "1820", "8950", "88.2%", "36200000.00", "35800000.00", "36500000.00", "-400000.00", "82.5"],
        ["Orthopedic Surgery", "2410", "11200", "91.0%", "44800000.00", "47200000.00", "48500000.00", "2400000.00", "96.0"],
        ["Intensive Care (ICU)", "980", "4820", "96.4%", "58400000.00", "63800000.00", "65200000.00", "5400000.00", "118.0"],
        ["Pediatrics & Neonatal", "1420", "5600", "78.5%", "28500000.00", "27900000.00", "28200000.00", "-600000.00", "64.0"],
        ["Radiology & Imaging", "12500", "0", "N/A", "22000000.00", "21500000.00", "21800000.00", "-500000.00", "45.0"],
        ["Oncology & Chemotherapy", "3150", "7800", "89.0%", "41200000.00", "43600000.00", "44500000.00", "2400000.00", "78.0"],
    ]
    for r in rows:
        table_data.append(r)

    # Clean Tufte table styling: no vertical gridlines, muted headers, crisp line under header
    t = Table(table_data, hAlign='LEFT')
    t.setStyle(TableStyle([
        ('FONTNAME', (0, 0), (-1, 0), 'Helvetica-Bold'),
        ('FONTSIZE', (0, 0), (-1, -1), 9),
        ('TEXTCOLOR', (0, 0), (-1, 0), colors.HexColor('#0F172A')),
        ('BACKGROUND', (0, 0), (-1, 0), colors.HexColor('#F1F5F9')),
        ('LINEBELOW', (0, 0), (-1, 0), 1.2, colors.HexColor('#0F172A')),
        ('LINEBELOW', (0, -1), (-1, -1), 1.2, colors.HexColor('#0F172A')),
        ('LINEBELOW', (0, 1), (-1, -2), 0.5, colors.HexColor('#E2E8F0')),
        ('ALIGN', (0, 0), (0, -1), 'LEFT'),
        ('ALIGN', (1, 0), (-1, -1), 'RIGHT'),
        ('VALIGN', (0, 0), (-1, -1), 'MIDDLE'),
        ('TOPPADDING', (0, 0), (-1, -1), 5),
        ('BOTTOMPADDING', (0, 0), (-1, -1), 5),
    ]))
    story.append(t)
    doc.build(story)
    return output_path


def generate_consulting_txt(output_path: Path) -> Path:
    headers = [
        "engagement_id", "client_name", "practice_area", 
        "budget_fees_nok", "actual_billing_nok", "forecast_billing_nok", 
        "incurred_hours", "target_hours", "realization_rate_pct", "lead_partner"
    ]
    rows = [
        ["ENG-2024-01", "Kongsberg Maritime", "Systems Optimization", "4500000.00", "4820000.00", "5100000.00", "2450.0", "2200.0", "96.5", "J. Hansen"],
        ["ENG-2024-02", "Statkraft Hydro", "Digital Transformation", "3200000.00", "3100000.00", "3300000.00", "1620.0", "1600.0", "98.2", "E. Solberg"],
        ["ENG-2024-03", "Yara International", "Supply Chain Decarbonization", "6800000.00", "7250000.00", "7500000.00", "3800.0", "3500.0", "94.8", "K. Moe"],
        ["ENG-2024-04", "DNB Markets", "Risk Governance Automation", "5100000.00", "4950000.00", "5200000.00", "2750.0", "2800.0", "99.1", "T. Olsen"],
        ["ENG-2024-05", "Equinor Renewables", "Offshore Wind Financial Model", "5900000.00", "6400000.00", "6600000.00", "3100.0", "2900.0", "95.0", "M. Lund"],
    ]
    output_path.parent.mkdir(parents=True, exist_ok=True)
    with output_path.open("w", newline="", encoding="utf-8") as f:
        writer = csv.writer(f, delimiter=",")
        writer.writerow(headers)
        writer.writerows(rows)
    return output_path


def generate_all_mock_data(base_dir: Path | None = None) -> dict[str, Path]:
    if base_dir is None:
        base_dir = Path(__file__).resolve().parents[1] / "test_data" / "unseen_businesses"
    base_dir.mkdir(parents=True, exist_ok=True)

    files = {
        "maritime_defense_vessel_excel": generate_maritime_defense_excel(base_dir / "maritime_defense_vessel.xlsx"),
        "saas_subscription_tsv": generate_saas_tsv(base_dir / "saas_subscription_platform.tsv"),
        "nordic_retail_semicolon_csv": generate_nordic_retail_semicolon_csv(base_dir / "nordic_retail_inventory.csv"),
        "offshore_marine_psv": generate_offshore_marine_psv(base_dir / "offshore_marine_drilling.psv"),
        "hospital_kpis_pdf": generate_hospital_kpi_pdf(base_dir / "hospital_executive_kpis.pdf"),
        "consulting_engagements_txt": generate_consulting_txt(base_dir / "management_consulting_engagements.txt"),
    }
    return files


if __name__ == "__main__":
    generated = generate_all_mock_data()
    print("Generated unseen mock datasets across multiple businesses and formats:")
    for name, path in generated.items():
        print(f" - {name}: {path} ({path.stat().st_size:,} bytes)")
